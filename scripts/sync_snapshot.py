#!/usr/bin/env python3
"""Build the static workbench snapshot from existing promotion outputs."""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import tempfile
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


TEAMS = [
    "一大队",
    "二大队",
    "三大队",
    "四大队",
    "五大队",
    "六大队",
    "七大队",
    "八大队",
    "九大队",
    "十大队",
    "会理大队",
    "德昌大队",
]
TEAM_ORDER = {team: index for index, team in enumerate(TEAMS)}
PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = PROJECT_ROOT / "docs" / "data" / "snapshot.json"
DEFAULT_DAILY_ROOT = Path.home() / "Documents" / "Codex"
DEFAULT_INPUT_DIR = Path.home() / "Desktop" / "推广通报输入"
DEFAULT_LEDGER = Path.home() / "Desktop" / "推广通报数据" / "推广汇总统计.xlsx"
REQUIRED_EXPORT_FILES = ("dailyA.tsv", "dailyB.tsv", "records_all.tsv")


class SnapshotError(RuntimeError):
    """Raised when source data cannot produce a safe snapshot."""


def normalize_employee_id(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(int(value))
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    raw = str(value).strip()
    if re.fullmatch(r"\d+\.0+", raw):
        return raw.split(".", 1)[0]
    return raw


def normalize_name(value: Any) -> str:
    return "" if value is None else str(value).strip()


def iso_date(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    raw = str(value).strip()
    if not raw:
        return ""
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return dt.datetime.strptime(raw[:10], pattern).date().isoformat()
        except ValueError:
            continue
    match = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", raw)
    if match:
        year, month, day = (int(part) for part in match.groups())
        try:
            return dt.date(year, month, day).isoformat()
        except ValueError:
            return ""
    return ""


def normalize_flag(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    try:
        numeric = int(float(value))
    except (TypeError, ValueError):
        return None
    return numeric if numeric in (0, 1) else None


def cell(row: Iterable[Any], index: int | None) -> Any:
    if index is None:
        return None
    values = list(row)
    return values[index] if 0 <= index < len(values) else None


def parse_roster_lines(
    sources: Iterable[tuple[str, Iterable[str]]],
) -> list[dict[str, Any]]:
    roster: dict[str, dict[str, str]] = {}
    for source_name, lines in sources:
        for line_number, line in enumerate(lines, 1):
            parts = line.rstrip("\n").split("|")
            if len(parts) < 4:
                continue
            _, raw_name, raw_id, raw_team = parts[:4]
            name = normalize_name(raw_name)
            employee_id = normalize_employee_id(raw_id)
            team = normalize_name(raw_team)
            if (
                not name
                or not employee_id
                or name in {"司机姓名", "#N/A"}
                or employee_id in {"司机工号", "#N/A"}
                or team not in TEAM_ORDER
            ):
                continue
            record = {"employeeId": employee_id, "name": name, "team": team}
            prior = roster.get(employee_id)
            if prior and prior != record:
                raise SnapshotError(
                    f"{source_name}:{line_number} 工号 {employee_id} 的姓名或队伍冲突"
                )
            roster[employee_id] = record

    return sorted(
        roster.values(),
        key=lambda row: (TEAM_ORDER[row["team"]], row["name"], row["employeeId"]),
    )


def read_roster(daily_dir: Path) -> list[dict[str, Any]]:
    sources = []
    for filename in ("dailyA.tsv", "dailyB.tsv"):
        path = daily_dir / filename
        if not path.exists():
            raise SnapshotError(f"缺少名单文件：{path}")
        sources.append((filename, path.read_text(encoding="utf-8").splitlines()))
    return parse_roster_lines(sources)


def summarize_ledger_rows(
    rows: Iterable[Iterable[Any]],
    month: str,
    valid_ids: set[str],
) -> dict[str, dict[str, int]]:
    out_dates: dict[str, set[str]] = defaultdict(set)
    miss_dates: dict[str, set[str]] = defaultdict(set)

    for row in rows:
        values = list(row)
        if len(values) < 5 or values[0] is None:
            continue
        date_value = iso_date(values[0])
        employee_id = normalize_employee_id(values[3])
        if (
            not date_value.startswith(month)
            or not employee_id
            or employee_id not in valid_ids
        ):
            continue
        out_dates[employee_id].add(date_value)
        if normalize_name(values[4]) == "否":
            miss_dates[employee_id].add(date_value)

    return {
        employee_id: {
            "outDays": len(dates),
            "misses": len(miss_dates.get(employee_id, set())),
        }
        for employee_id, dates in out_dates.items()
    }


def read_ledger_metrics(
    ledger_path: Path, month: str, valid_ids: set[str]
) -> dict[str, dict[str, int]]:
    if not ledger_path.exists():
        raise SnapshotError(f"缺少个人台账：{ledger_path}")
    workbook = load_workbook(ledger_path, read_only=True, data_only=True)
    try:
        if "个人台账" not in workbook.sheetnames:
            raise SnapshotError("个人台账缺少「个人台账」工作表")
        return summarize_ledger_rows(
            workbook["个人台账"].iter_rows(min_row=2, values_only=True),
            month,
            valid_ids,
        )
    finally:
        workbook.close()


def count_monthly_submissions(records_path: Path, month: str) -> dict[str, int]:
    if not records_path.exists():
        raise SnapshotError(f"缺少推广提交记录：{records_path}")
    counts: dict[str, int] = defaultdict(int)
    for line in records_path.read_text(encoding="utf-8").splitlines():
        parts = line.split("|")
        if len(parts) < 3:
            continue
        name = normalize_name(parts[1])
        try:
            timestamp = dt.datetime.strptime(parts[2].strip(), "%Y/%m/%d %H:%M")
        except ValueError:
            continue
        if name and timestamp.strftime("%Y-%m") == month:
            counts[name] += 1
    return dict(counts)


def find_header_index(headers: list[str], *names: str) -> int | None:
    for name in names:
        if name in headers:
            return headers.index(name)
    return None


def parse_promo_rows(
    headers: list[str],
    rows: Iterable[Iterable[Any]],
    source_date: str,
) -> tuple[dict[str, int], list[dict[str, Any]]]:
    employee_index = find_header_index(headers, "推荐司机工号")
    bd_type_index = find_header_index(headers, "bd_type")
    date_index = find_header_index(headers, "日期")
    if employee_index is None or bd_type_index is None or date_index is None:
        raise SnapshotError("当月推广数据缺少推荐司机工号、bd_type 或日期列")

    # The daily pipeline historically falls back to column 17 when the
    # explicit true-promotion column is absent.
    true_promotion_index = find_header_index(headers, "是否非真实推广")
    if true_promotion_index is None and len(headers) > 17:
        true_promotion_index = 17

    month = source_date[:7]
    transfer_counts: dict[str, int] = defaultdict(int)
    orders: list[dict[str, Any]] = []

    for row in rows:
        values = list(row)
        date_value = iso_date(cell(values, date_index))
        employee_id = normalize_employee_id(cell(values, employee_index))
        if (
            not employee_id
            or not date_value
            or not date_value.startswith(month)
            or date_value > source_date
        ):
            continue

        order = {
            "employeeId": employee_id,
            "date": date_value,
            "nonOffline": normalize_flag(cell(values, 17)),
            "abnormalScan": normalize_flag(cell(values, 18)),
            "burner": normalize_flag(cell(values, 19)),
            "regularCustomer": normalize_flag(cell(values, 20)),
            "cheating": normalize_flag(cell(values, 21)),
        }
        orders.append(order)

        bd_type = normalize_name(cell(values, bd_type_index))
        true_promotion_flag = normalize_flag(cell(values, true_promotion_index))
        if bd_type and true_promotion_flag == 0:
            transfer_counts[employee_id] += 1

    orders.sort(key=lambda item: (item["date"], item["employeeId"]), reverse=True)
    return dict(transfer_counts), orders


def read_promo_data(
    promo_path: Path, source_date: str
) -> tuple[dict[str, int], list[dict[str, Any]]]:
    if not promo_path.exists():
        raise SnapshotError(f"缺少当月推广数据：{promo_path}")
    workbook = load_workbook(promo_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        headers = [normalize_name(value) for value in next(rows)]
        return parse_promo_rows(headers, rows, source_date)
    finally:
        workbook.close()


def build_driver_rows(
    roster: list[dict[str, Any]],
    ledger_metrics: dict[str, dict[str, int]],
    submission_counts: dict[str, int],
    transfer_counts: dict[str, int],
) -> list[dict[str, Any]]:
    drivers = []
    for person in roster:
        employee_id = person["employeeId"]
        ledger = ledger_metrics.get(employee_id, {"outDays": 0, "misses": 0})
        total_promotions = submission_counts.get(person["name"], 0)
        drivers.append(
            {
                "employeeId": employee_id,
                "name": person["name"],
                "team": person["team"],
                "outDays": int(ledger["outDays"]),
                "realTransfers": int(transfer_counts.get(employee_id, 0)),
                "promotionDelta": total_promotions - int(ledger["misses"]),
            }
        )
    return drivers


def read_latest_sent(daily_dir: Path) -> dict[str, Any]:
    sent_files = sorted(
        daily_dir.glob(".sent-*.json"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    if not sent_files:
        raise SnapshotError(f"导出目录没有成功的发送记录：{daily_dir}")
    try:
        payload = json.loads(sent_files[0].read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise SnapshotError(f"无法读取发送记录：{sent_files[0]}") from error
    source_date = iso_date(payload.get("date"))
    if not source_date:
        raise SnapshotError(f"发送记录缺少有效日期：{sent_files[0]}")
    return {
        "sourceDate": source_date,
        "syncedAt": normalize_name(payload.get("sent_at")) or source_date,
    }


def discover_latest_daily_dir(root: Path) -> Path:
    candidates = {
        path.parent
        for path in root.glob("**/tongbao-*/dailyA.tsv")
        if all((path.parent / filename).exists() for filename in REQUIRED_EXPORT_FILES)
        and list(path.parent.glob(".sent-*.json"))
    }
    if not candidates:
        raise SnapshotError(f"没有找到完整且已成功发送的每日导出：{root}")
    return max(candidates, key=lambda path: path.stat().st_mtime)


def discover_latest_promo(input_dir: Path, source_date: str) -> Path:
    month_number = int(source_date[5:7])
    exact_name = f"{month_number}月推广数据.xlsx"
    exact_path = input_dir / exact_name
    if exact_path.exists():
        return exact_path
    candidates = [
        path
        for path in input_dir.glob("*月推广数据.xlsx")
        if not path.name.startswith("~$")
    ]
    if not candidates:
        raise SnapshotError(f"没有找到当月推广数据：{input_dir}")
    return max(candidates, key=lambda path: path.stat().st_mtime)


def read_bt5_metrics(daily_dir: Path) -> dict[tuple[str, str], dict[str, int]]:
    metrics: dict[tuple[str, str], dict[str, int]] = {}
    for path in sorted(daily_dir.glob("bt5_*.tsv")):
        team = path.stem.removeprefix("bt5_")
        if team not in TEAM_ORDER:
            continue
        for line in path.read_text(encoding="utf-8").splitlines():
            parts = line.split("\t")
            if len(parts) < 5:
                continue
            name = normalize_name(parts[0])
            try:
                metrics[(team, name)] = {
                    "realTransfers": int(parts[2]),
                    "promotionDelta": int(parts[4]),
                }
            except ValueError as error:
                raise SnapshotError(f"bt5 数值无效：{path} / {name}") from error
    return metrics


def validate_against_bt5(
    drivers: list[dict[str, Any]],
    bt5_metrics: dict[tuple[str, str], dict[str, int]],
) -> None:
    mismatches = []
    for driver in drivers:
        expected = bt5_metrics.get((driver["team"], driver["name"]))
        if not expected:
            continue
        if (
            driver["realTransfers"] != expected["realTransfers"]
            or driver["promotionDelta"] != expected["promotionDelta"]
        ):
            mismatches.append(
                f'{driver["team"]}/{driver["name"]}: '
                f'转单 {driver["realTransfers"]}!={expected["realTransfers"]}, '
                f'差值 {driver["promotionDelta"]}!={expected["promotionDelta"]}'
            )
    if mismatches:
        sample = "；".join(mismatches[:5])
        raise SnapshotError(f"页面数据与最新 bt5 对账失败：{sample}")


def validate_snapshot(snapshot: dict[str, Any]) -> None:
    meta = snapshot.get("meta", {})
    drivers = snapshot.get("drivers", [])
    orders = snapshot.get("orders", [])
    if not meta.get("periodStart") or not meta.get("periodEnd"):
        raise SnapshotError("快照缺少统计周期")
    if not drivers:
        raise SnapshotError("快照没有司机数据")

    seen_ids: set[str] = set()
    team_counts: dict[str, int] = defaultdict(int)
    for driver in drivers:
        employee_id = normalize_employee_id(driver.get("employeeId"))
        team = normalize_name(driver.get("team"))
        if not employee_id or not normalize_name(driver.get("name")):
            raise SnapshotError("司机姓名或工号为空")
        if employee_id in seen_ids:
            raise SnapshotError(f"司机工号重复：{employee_id}")
        if team not in TEAM_ORDER:
            raise SnapshotError(f"司机队伍无效：{team}")
        seen_ids.add(employee_id)
        team_counts[team] += 1

    missing_teams = [team for team in TEAMS if team_counts[team] == 0]
    if missing_teams:
        raise SnapshotError(f"以下队伍没有司机：{', '.join(missing_teams)}")

    previous_key = ""
    for order in orders:
        employee_id = normalize_employee_id(order.get("employeeId"))
        date_value = iso_date(order.get("date"))
        if not employee_id or not date_value:
            raise SnapshotError("订单工号或日期无效")
        for field in (
            "nonOffline",
            "abnormalScan",
            "burner",
            "regularCustomer",
            "cheating",
        ):
            if order.get(field) not in (0, 1, None):
                raise SnapshotError(f"订单字段 {field} 无效")
        sort_key = f"{date_value}:{employee_id}"
        if previous_key and sort_key > previous_key:
            raise SnapshotError("订单未按日期倒序排列")
        previous_key = sort_key


def build_snapshot(
    *,
    daily_dir: Path,
    promo_path: Path,
    ledger_path: Path,
) -> dict[str, Any]:
    sent = read_latest_sent(daily_dir)
    source_date = sent["sourceDate"]
    month = source_date[:7]
    period_start = f"{month}-01"

    roster = read_roster(daily_dir)
    valid_ids = {row["employeeId"] for row in roster}
    ledger_metrics = read_ledger_metrics(ledger_path, month, valid_ids)
    submission_counts = count_monthly_submissions(
        daily_dir / "records_all.tsv", month
    )
    transfer_counts, orders = read_promo_data(promo_path, source_date)
    drivers = build_driver_rows(
        roster,
        ledger_metrics,
        submission_counts,
        transfer_counts,
    )
    validate_against_bt5(drivers, read_bt5_metrics(daily_dir))

    snapshot = {
        "meta": {
            "schemaVersion": 1,
            "periodStart": period_start,
            "periodEnd": source_date,
            "sourceDate": source_date,
            "syncedAt": sent["syncedAt"],
            "driverCount": len(drivers),
            "orderCount": len(orders),
            "teamCount": len(TEAMS),
        },
        "drivers": drivers,
        "orders": orders,
    }
    validate_snapshot(snapshot)
    return snapshot


def write_snapshot(snapshot: dict[str, Any], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n"
    with tempfile.NamedTemporaryFile(
        "w",
        encoding="utf-8",
        dir=output_path.parent,
        prefix=".snapshot.",
        suffix=".tmp",
        delete=False,
    ) as temporary:
        temporary.write(payload)
        temporary_path = Path(temporary.name)
    os.replace(temporary_path, output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--daily-dir", type=Path)
    parser.add_argument("--daily-root", type=Path, default=DEFAULT_DAILY_ROOT)
    parser.add_argument("--promo", type=Path)
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--ledger", type=Path, default=DEFAULT_LEDGER)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        daily_dir = args.daily_dir or discover_latest_daily_dir(args.daily_root)
        sent = read_latest_sent(daily_dir)
        promo_path = args.promo or discover_latest_promo(
            args.input_dir, sent["sourceDate"]
        )
        snapshot = build_snapshot(
            daily_dir=daily_dir,
            promo_path=promo_path,
            ledger_path=args.ledger,
        )
        write_snapshot(snapshot, args.output)
    except SnapshotError as error:
        print(f"同步失败：{error}")
        return 1

    print(
        json.dumps(
            {
                "ok": True,
                "dailyDir": str(daily_dir),
                "promo": str(promo_path),
                "output": str(args.output),
                "sourceDate": snapshot["meta"]["sourceDate"],
                "drivers": snapshot["meta"]["driverCount"],
                "orders": snapshot["meta"]["orderCount"],
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

